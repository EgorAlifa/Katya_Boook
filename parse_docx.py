"""
Single-pass docx -> classified.json parser, using python-docx directly against
the real document object model (paragraphs, runs, tables, drawings) instead of
round-tripping through pandoc markdown + regex.

Why: pandoc's markdown text representation threw away structural signal we
actually need (real bold/italic/subscript run flags collapse into `**`/`*`/`~..~`
markdown escaping, which is lossy and fragile — this session hit real content
bugs from it: subscript tildes leaking into prose, an image glued mid-sentence
silently splitting a sentence, images glued directly to their own caption with
no blank line getting entirely discarded by a naive "image-only block" length
heuristic, an inline image's exact position in a run of text being unrecoverable
from flattened markdown so it had to be *guessed*). Reading the OOXML directly
sidesteps all of that: run order, bold/italic/subscript, and each drawing's
exact position in the run sequence are all just *there*, unambiguously — no
markdown round-trip, no guessing, nothing to escape.

Also recovers content the old pandoc-based pipeline silently dropped outright:
parse_book.py sliced the source at the "СОДЕРЖАНИЕ" heading and never looked at
anything after it — which is where the book's own colophon lives (real print
format "60x84/8", real "Тираж/Заказ", etc.), previously replaced by hand-typed
placeholder text in build-full.js. See COLOPHON_LINES below and README.md.

Replaces: pandoc's markdown export, parse_book.py, extract_image_layout.py,
and classify_book.py's markdown-regex classification. Writes classified.json
in the exact same block schema build-full.js already consumes (build-full.js
and design.js are UNCHANGED), plus img/*  and colophon_lines.json.
"""
import os
import re
import json
import hashlib
from collections import Counter

import docx
from docx.oxml.ns import qn

SRC_DOCX = "вар22223333.docx"
IMG_DIR = "img"
EMU_PER_INCH = 914400

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def qn2(tag):
    prefix, local = tag.split(":")
    return "{%s}%s" % (NS[prefix], local)


def despace(s):
    return s.replace(" ", "").strip()


os.makedirs(IMG_DIR, exist_ok=True)
_written_images = {}  # filename -> sha1, so a byte-identical re-embed is a silent no-op
doc = docx.Document(SRC_DOCX)


def extract_drawing(drawing_el, part):
    """Real display size (wp:extent) + rotation (a:xfrm/@rot) + the image's own
    bytes, straight from the OOXML — exactly what Word itself used, not a
    pixel-derived guess. Writes the blob to img/<original filename> (same
    naming convention the old unzip step used, so design.js needs no changes)."""
    extent = drawing_el.find(".//wp:extent", NS)
    blip = drawing_el.find(".//a:blip", NS)
    if extent is None or blip is None:
        return None
    rId = blip.get(qn2("r:embed"))
    if not rId or rId not in part.related_parts:
        return None
    image_part = part.related_parts[rId]
    filename = os.path.basename(image_part.partname)
    blob = image_part.blob
    sha1 = hashlib.sha1(blob).hexdigest()
    prev = _written_images.get(filename)
    if prev is None:
        with open(os.path.join(IMG_DIR, filename), "wb") as f:
            f.write(blob)
        _written_images[filename] = sha1
    elif prev != sha1:
        print(f"WARNING: {filename} written with two different contents (rId {rId})")

    w_in = int(extent.get("cx")) / EMU_PER_INCH
    h_in = int(extent.get("cy")) / EMU_PER_INCH
    rotation = 0.0
    xfrm = drawing_el.find(".//pic:spPr/a:xfrm", NS)
    if xfrm is not None:
        rot = xfrm.get("rot")
        if rot:
            rotation = int(rot) / 60000.0
    return {"file": filename, "w": round(w_in, 4), "h": round(h_in, 4), "rotation": rotation}


def paragraph_pieces(p):
    """Ordered list of ('text', text, bold, subscript) / ('image', info) pieces,
    in REAL document order — an inline image's exact position relative to the
    surrounding text is read straight off the run sequence, never inferred."""
    pieces = []
    part = p.part
    for r in p.runs:
        drawings = r._r.findall(".//" + qn2("w:drawing"))
        for dr in drawings:
            info = extract_drawing(dr, part)
            if info:
                pieces.append(("image", info))
        if r.text:
            pieces.append(("text", r.text, bool(r.bold), bool(r.font.subscript)))
    return pieces


def full_text(pieces):
    return despace("".join(t[1] for t in pieces if t[0] == "text"))


def all_bold(pieces):
    tp = [t for t in pieces if t[0] == "text" and t[1].strip()]
    return bool(tp) and all(t[2] for t in tp)


# ---- generic interleaver: turns a piece list into ordered para/image blocks,
# preserving true position (consecutive images bundle into one side-by-side
# block; text between/around them becomes its own paragraph) -----------------
def interleave_blocks(pieces):
    out = []
    text_buf = []
    img_buf = []

    def flush_text():
        if text_buf:
            t = despace("".join(text_buf))
            if t:
                out.append({"type": "para", "text": t})
            text_buf.clear()

    def flush_img():
        if img_buf:
            out.append({"type": "image", "images": list(img_buf)})
            img_buf.clear()

    for kind, *rest in pieces:
        if kind == "image":
            flush_text()
            img_buf.append(rest[0])
        else:
            flush_img()
            text_buf.append(rest[0])
    flush_text()
    flush_img()
    return out


CAPTION_RE = re.compile(r"^Рисунок\s+([\w.\-]+)\.?\s*$")


def try_caption(pieces):
    """Bold 'Рисунок N-M.' run(s) followed by italic description — possibly more
    than one figure captioned in the same paragraph. Returns None if this isn't
    a caption paragraph at all."""
    if not despace(full_text(pieces)).startswith("Рисунок"):
        return None
    items = []
    cur_num, cur_desc = None, []
    for piece in pieces:
        if piece[0] != "text":
            continue
        _, txt, bold, _ = piece
        if bold and re.match(r"^\s*Рисунок\b", txt):
            if cur_num:
                items.append({"num": cur_num.strip(), "desc": despace("".join(cur_desc)).strip(" *")})
            cur_num, cur_desc = txt, []
        elif cur_num is not None:
            cur_desc.append(txt)
    if cur_num:
        items.append({"num": despace(cur_num).strip(), "desc": despace("".join(cur_desc)).strip(" *")})
    return items or None


# "1/S=1/T<sub>Земли</sub> – 1/T<sub>внеш</sub>." — driven by the REAL subscript
# run flag (w:vertAlign="subscript"), not a markdown-tilde guess.
FORMULA_RE = re.compile(r"1/S\s*=\s*1/T\x01([^\x02]*)\x02\s*[–\-—−]\s*1/T\x01([^\x02]*)\x02\.?")


def try_formula(pieces):
    # Word often splits one subscripted word across several adjacent runs
    # (e.g. "Земли" as separate 'Земл' + 'и' runs, both subscript=True) — merge
    # consecutive subscript runs into ONE \x01..\x02 span, don't wrap each run
    # individually, or the marker pair closes mid-word and the regex below
    # either misses the formula entirely or truncates the subscript it captures.
    marked = []
    open_sub = False
    for piece in pieces:
        if piece[0] != "text":
            continue
        _, txt, bold, sub = piece
        if sub and not open_sub:
            marked.append("\x01")
            open_sub = True
        elif not sub and open_sub:
            marked.append("\x02")
            open_sub = False
        marked.append(txt)
    if open_sub:
        marked.append("\x02")
    marked_text = despace("".join(marked))
    m = FORMULA_RE.search(marked_text)
    if not m:
        return None
    pre = despace(marked_text[: m.start()].replace("\x01", "").replace("\x02", ""))
    post = despace(marked_text[m.end() :].replace("\x01", "").replace("\x02", ""))
    blocks = []
    if pre:
        blocks.append({"type": "para", "text": pre})
    blocks.append({"type": "formula", "den1": despace(m.group(1)), "den2": despace(m.group(2))})
    if post:
        blocks.append({"type": "para", "text": post})
    return blocks


def pull_images(pieces):
    """All image pieces anywhere in the paragraph, as one bundled image block
    (or []). Word labels like ГЛАВА/ЗАДАЧА/ОТВЕТ are occasionally glued to a
    figure in the very same paragraph (e.g. two images immediately ahead of
    "ОТВЕТ" in one run sequence) — every "whole paragraph is one label" branch
    below must check this, or the figure silently vanishes."""
    imgs = [t[1] for t in pieces if t[0] == "image"]
    return [{"type": "image", "images": imgs}] if imgs else []


def classify_paragraph(p):
    pieces = paragraph_pieces(p)
    if not pieces:
        return []
    text_pieces = [t for t in pieces if t[0] == "text"]
    ft = full_text(pieces)

    if not text_pieces:
        return interleave_blocks(pieces)  # pure image(s)

    # a couple of paragraphs in the source are a single stray "." with nothing
    # else — a leftover editing artifact (deleted text, punctuation left
    # behind), not real content. Drop only bare punctuation, never real text.
    if re.fullmatch(r"[.\-–—,;:]+", ft) and not any(t[0] == "image" for t in pieces):
        return []

    if re.match(r"^ГЛАВА\s*\d+$", ft):
        return pull_images(pieces) + [{"type": "chapter_mark", "text": ft}]
    if ft == "ЗАКЛЮЧЕНИЕ":
        return pull_images(pieces) + [{"type": "conclusion_mark"}]
    m = re.match(r"^ЗАДАЧА\s*(\d+)\s*-\s*(\d+)$", ft)
    if m:
        return pull_images(pieces) + [{"type": "task_label", "text": f"ЗАДАЧА {m.group(1)}-{m.group(2)}"}]
    if ft == "ОТВЕТ":
        return pull_images(pieces) + [{"type": "answer_label"}]

    fm = try_formula(pieces)
    if fm:
        return pull_images(pieces) + fm

    cap = try_caption(pieces)
    if cap:
        # the figure itself is often glued into the SAME paragraph as its own
        # caption (image run(s) then "Рисунок N-M." bold run then description)
        # rather than a separate paragraph before it.
        return pull_images(pieces) + [{"type": "caption", "items": cap}]

    if all_bold(pieces) and not any(t[0] == "image" for t in pieces) and len(ft) < 140:
        if ft.isupper() and len(ft) > 4:
            return [{"type": "chapter_title", "text": ft}]
        return [{"type": "subheading", "text": ft}]

    # drop-cap: first text piece is a single bold uppercase letter, the rest of
    # the paragraph isn't (fully) bold. Any images can land on EITHER side of
    # that letter (a figure glued ahead of the drop-cap paragraph, or one glued
    # mid-sentence further into it) — split around the letter's own run only,
    # and let interleave_blocks place images correctly on each side rather
    # than assuming "images always precede" or "always follow".
    first = text_pieces[0]
    first_stripped = despace(first[1]).strip()
    if len(first_stripped) == 1 and first[2] and first_stripped.isupper() and not all_bold(pieces):
        letter_idx = pieces.index(first)
        # if the letter run got split (e.g. by a following empty/space bold
        # run), fold those into the "letter" span too
        end_idx = letter_idx + 1
        while end_idx < len(pieces) and pieces[end_idx][0] == "text" and pieces[end_idx][2] and not despace(pieces[end_idx][1]).strip():
            end_idx += 1
        before, after = pieces[:letter_idx], pieces[end_idx:]
        out = interleave_blocks(before)
        after_blocks = interleave_blocks(after)
        if after_blocks and after_blocks[0]["type"] == "para":
            out.append({"type": "para_dropcap", "letter": first_stripped, "text": after_blocks[0]["text"]})
            out.extend(after_blocks[1:])
        else:
            out.append({"type": "para_dropcap", "letter": first_stripped, "text": ""})
            out.extend(after_blocks)
        return out

    # default paragraph. A figure genuinely glued to the FRONT of a paragraph
    # (nothing real before it) splits fine — interleave_blocks below already
    # gets that right, image block then its own following text. But when real
    # prose precedes the first image — the image is just anchored somewhere
    # inside a sentence, a Word inline-anchor artifact, not a meaningful split
    # point — interleaving verbatim visibly breaks the sentence in two around
    # it (confirmed live: "...длится 686,98 земных [image] суток." coming out
    # as two separate paragraphs either side of the figure, its caption then
    # reading like it belongs to the SECOND half). Keep the sentence whole and
    # put every image from the paragraph after it instead.
    if pieces[0][0] != "image" and any(t[0] == "image" for t in pieces):
        merged_text = despace("".join(t[1] for t in pieces if t[0] == "text"))
        imgs = [t[1] for t in pieces if t[0] == "image"]
        out = []
        if merged_text:
            out.append({"type": "para", "text": merged_text})
        out.append({"type": "image", "images": imgs})
        return out

    return interleave_blocks(pieces)


# ---- callout tables ("Это интересно") and the planet-drift data table ------
def iter_unique_cells(table):
    seen = set()
    for row in table.rows:
        for cell in row.cells:
            key = id(cell._tc)
            if key in seen:
                continue
            seen.add(key)
            yield cell


def classify_table(table):
    cells = list(iter_unique_cells(table))
    if not cells:
        return None
    first_cell_text = despace(cells[0].paragraphs[0].text) if cells[0].paragraphs else ""
    if despace(first_cell_text).upper().replace(" ", "") in ("ЭТОИНТЕРЕСНО",) or "это интересно" in first_cell_text.lower():
        label_seen = False
        title = None
        body_parts = []
        images = []
        for cell in cells:
            for p in cell.paragraphs:
                pieces = paragraph_pieces(p)
                for piece in pieces:
                    if piece[0] == "image":
                        images.append(piece[1])
                text_pieces = [t for t in pieces if t[0] == "text"]
                if not text_pieces:
                    continue
                txt = full_text(pieces)
                if not txt:
                    continue
                if not label_seen and "это интересно" in despace(txt).lower():
                    label_seen = True
                    continue
                if title is None and all_bold(pieces):
                    title = txt
                    continue
                body_parts.append(txt)
        return {"type": "callout", "title": title, "body_text": " ".join(body_parts).strip(), "images": images}

    # planet angular-drift table: 4 rows x 9 cols of numbers — rendered by
    # build-full.js's own renderPlanetTable() (verified to match this table's
    # real data), so the parser only needs to flag its presence.
    if len(table.rows) == 4 and len(table.columns) == 9:
        return {"type": "dashtable"}

    print("WARNING: unrecognized table shape", len(table.rows), "x", len(table.columns), "- skipped")
    return None


# ---- locate the body range: from "ВВЕДЕНИЕ" (inclusive) to the manuscript's
# own "СОДЕРЖАНИЕ" heading (exclusive — it's a hand-typed ToC with no page
# numbers, redundant with our real Word TOC field; see README) -------------
def find_para_index(target):
    for i, p in enumerate(doc.paragraphs):
        if despace(p.text) == target:
            return i
    raise ValueError(f"paragraph {target!r} not found")


intro_idx = find_para_index("ВВЕДЕНИЕ")
toc_idx = find_para_index("СОДЕРЖАНИЕ")

body_elements = list(doc.element.body.iterchildren())
# map each body child element to a (kind, index-in-doc.paragraphs/doc.tables) pair
p_iter = iter(enumerate(doc.paragraphs))
t_iter = iter(enumerate(doc.tables))
ordered = []
for el in body_elements:
    if el.tag == qn("w:p"):
        idx, para = next(p_iter)
        ordered.append(("p", idx, para))
    elif el.tag == qn("w:tbl"):
        idx, table = next(t_iter)
        ordered.append(("t", idx, table))

# slice to the body range by paragraph index
start_pos = next(i for i, (k, idx, _) in enumerate(ordered) if k == "p" and idx == intro_idx)
end_pos = next(i for i, (k, idx, _) in enumerate(ordered) if k == "p" and idx == toc_idx)
body_ordered = ordered[start_pos:end_pos]

out = []
for kind, idx, obj in body_ordered:
    if kind == "p":
        out.extend(classify_paragraph(obj))
    else:
        c = classify_table(obj)
        if c:
            out.append(c)

json.dump(out, open("classified.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("blocks:", len(out))
print(Counter(b["type"] for b in out))

# ---- real colophon text (the book's own tail, from the mini title-page after
# the manuscript's hand-typed TOC through to the end) — replaces the
# previously hand-typed placeholder text in build-full.js -------------------
colophon_start = find_para_index("Заключение") + 1
# the FIRST "Заключение" match is the actual chapter heading inside the body
# range already consumed above; find_para_index returns the first occurrence
# overall, so re-search from just after toc_idx for the second (tail) one.
for i in range(toc_idx, len(doc.paragraphs)):
    if despace(doc.paragraphs[i].text) == "Заключение":
        colophon_start = i + 1
        break
colophon_lines = [despace(doc.paragraphs[i].text) for i in range(colophon_start, len(doc.paragraphs))]
colophon_lines = [l for l in colophon_lines if l]
json.dump(colophon_lines, open("colophon_lines.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("colophon lines:", colophon_lines)
